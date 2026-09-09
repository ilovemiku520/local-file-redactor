"""Bounded parsers and clean document reconstruction."""
from __future__ import annotations
import csv
from io import BytesIO, StringIO
import math
import os
from pathlib import Path
import re
import zipfile
from PIL import Image, ImageOps, ImageDraw, ImageFont
from lxml import etree
from .storage import ROOT, UserError

FORMATS = {'txt':['txt'],'csv':['csv'],'docx':['docx','pdf'],'xlsx':['xlsx','pdf'],
           'pdf':['pdf'], **{x:['png','jpg','jpeg','webp','bmp'] for x in ('png','jpg','jpeg','webp','bmp')}}
IMAGE_FORMATS = {'PNG','JPEG','WEBP','BMP'}
MAX_PIXELS = 40_000_000
Image.MAX_IMAGE_PIXELS = MAX_PIXELS


def text_decode(data):
    if len(data)>8_000_000:
        raise UserError('文本/CSV 超过 800 万字节，请拆分后处理。')
    if data.startswith((b'\xff\xfe', b'\xfe\xff')):
        return data.decode('utf-16'), 'UTF-16'
    for encoding in ('utf-8-sig','gb18030'):
        try:
            text=data.decode(encoding)
            if '\x00' in text:
                raise UserError('文件包含二进制内容，不能作为文本处理。')
            return text, encoding
        except UnicodeError:
            continue
    raise UserError('无法识别文本编码，请先另存为 UTF-8。')


def load_image(data):
    with Image.open(BytesIO(data)) as source:
        if source.format not in IMAGE_FORMATS or getattr(source,'n_frames',1) != 1:
            raise UserError('仅支持单帧 PNG、JPEG、WebP、BMP 图片。')
        if source.width * source.height > MAX_PIXELS:
            raise UserError('图片超过 4000 万像素上限。')
        fixed = ImageOps.exif_transpose(source).convert('RGBA')
        background = Image.new('RGBA', fixed.size, 'white')
        background.alpha_composite(fixed)
        return background.convert('RGB')


def image_bytes(image, extension='png'):
    out=BytesIO()
    options={'quality':95,'subsampling':0} if extension in ('jpg','jpeg') else ({'lossless':True} if extension=='webp' else {})
    image.convert('RGB').save(out,format={'jpg':'JPEG','jpeg':'JPEG'}.get(extension,extension.upper()),**options)
    return out.getvalue()


def package(data, extension):
    if not data.startswith(b'PK'):
        raise UserError('文件内容与 Office 扩展名不匹配。')
    try:
        archive=zipfile.ZipFile(BytesIO(data))
        members=archive.infolist()
        if len(members)>10000 or sum(i.file_size for i in members)>512*1024**2:
            raise UserError('Office 文件解压规模超过限制。')
        for item in members:
            if item.filename.startswith(('/', '\\')) or '..' in Path(item.filename).parts or ':' in item.filename:
                raise UserError('Office 文件包含异常路径。')
            if item.file_size>1024**2 and item.file_size/max(item.compress_size,1)>500:
                raise UserError('Office 文件压缩比异常。')
            lower=item.filename.lower()
            if any(x in lower for x in ('vbaproject','/embeddings/','/activex/','/externallinks/')):
                raise UserError('文件含宏、嵌入对象或外部数据，请先另存为静态文档。')
            if lower.endswith(('.xml','.rels')):
                content=archive.read(item)
                if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
                    raise UserError('XML 文件包含不允许的实体声明。')
                etree.fromstring(content,parser=etree.XMLParser(resolve_entities=False,no_network=True))
        expected='word/document.xml' if extension=='docx' else 'xl/workbook.xml'
        if expected not in archive.namelist():
            raise UserError('Office 文件类型不匹配。')
        return archive
    except (zipfile.BadZipFile,etree.XMLSyntaxError) as exc:
        raise UserError('Office 文件损坏或格式不受支持。') from exc


def csv_safe(value):
    if re.fullmatch(r'\s*-?\d+(?:\.\d+)?\s*',value):
        return value,False
    candidate=value.lstrip(' \r\n\t')
    if candidate.startswith(('=','+','-','@')) or value.startswith(('\t','\r','\n')):
        return "'"+value,True
    return value,False


def parse(data, extension, add_image):
    blocks,layout,warnings=[],[],[]
    def block(text, label='', context=''):
        index=len(blocks)
        blocks.append({'id':f'b{index}','text':str(text),'label':label,'context':context})
        return f'b{index}'
    if extension=='txt':
        text,encoding=text_decode(data)
        # Preserve exact separators while keeping each source block manageable.
        parts=re.findall(r'[^\r\n]*(?:\r\n|\r|\n)|[^\r\n]+$',text) or ['']
        for index,part in enumerate(parts):
            layout.append({'kind':'p','block':block(part,f'段落 {index+1}')})
        return {'blocks':blocks,'layout':layout,'warnings':warnings,'encoding':encoding}
    if extension=='csv':
        text,encoding=text_decode(data)
        try:
            dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t|')
            delimiter=dialect.delimiter
        except csv.Error:
            delimiter=','
        rows=list(csv.reader(StringIO(text,newline=''),delimiter=delimiter,strict=True))
        if sum(len(row) for row in rows)>200000:
            raise UserError('CSV 超过 20 万单元格上限。')
        refs=[]
        for r,row in enumerate(rows):
            refs.append([block(value,f'第 {r+1} 行 / 第 {c+1} 列',rows[0][c] if r and c<len(rows[0]) else '') for c,value in enumerate(row)])
        return {'blocks':blocks,'layout':refs,'warnings':warnings,'delimiter':delimiter,'encoding':encoding}
    if extension=='docx':
        from docx import Document
        from docx.oxml.ns import qn
        archive=package(data,extension)
        names=archive.namelist()
        for name in names:
            if name.startswith('word/charts/') or name.startswith('word/diagrams/'):
                raise UserError('Word 含图表或 SmartArt，请先转换为 PDF 后处理。')
        document=Document(BytesIO(data))
        # Explicit final-content policy: deleted revisions and hidden runs are excluded.
        for part in document.part.package.parts:
            root=getattr(part,'element',getattr(part,'_element',None))
            if root is not None:
                for element in list(root.iter(qn('w:del'))):
                    element.getparent().remove(element)
                for run in list(root.iter(qn('w:r'))):
                    if run.find('.//'+qn('w:vanish')) is not None:
                        run.getparent().remove(run)
                if any(list(root.iter(qn(tag))) for tag in ('w:txbxContent','m:oMath','w:altChunk','w:object')):
                    raise UserError('Word 含文本框或公式对象，请先转换为 PDF 后处理。')
        def nodes(parent,part):
            output=[]
            for element in parent:
                if element.tag==qn('w:p'):
                    text=''.join((n.text or '') if n.tag==qn('w:t') else ('\t' if n.tag==qn('w:tab') else '\n')
                                 for n in element.iter() if n.tag in (qn('w:t'),qn('w:tab'),qn('w:br')))
                    output.append({'kind':'p','block':block(text,f'段落 {len(blocks)+1}')})
                    for blip in element.iter(qn('a:blip')):
                        rid=blip.get(qn('r:embed'))
                        if not rid or rid not in part.related_parts:
                            raise UserError('Word 含外链图片，需先嵌入图片或转换为 PDF。')
                        image=load_image(part.related_parts[rid].blob)
                        output.append({'kind':'image','image':add_image(image)})
                elif element.tag==qn('w:tbl'):
                    rows=[]
                    for row in element.findall(qn('w:tr')):
                        rows.append([nodes(cell,part) for cell in row.findall(qn('w:tc'))])
                    output.append({'kind':'table','rows':rows})
                elif element.tag==qn('w:sdt'):
                    content=element.find(qn('w:sdtContent'))
                    if content is not None: output.extend(nodes(content,part))
                elif element.tag==qn('w:ins'):
                    output.extend(nodes(element,part))
            return output
        layout=nodes(document.element.body,document.part)
        seen=set()
        for section in document.sections:
            for field in ('header','first_page_header','even_page_header','footer','first_page_footer','even_page_footer'):
                part=getattr(section,field).part
                if part.partname not in seen:
                    seen.add(part.partname)
                    extra=nodes(part.element,part)
                    if any(blocks[int(n['block'][1:])]['text'].strip() if n['kind']=='p' else True for n in extra):
                        layout.extend(extra)
                        warnings.append('页眉/页脚内容已整理至正文末尾。')
        for name in ('word/footnotes.xml','word/endnotes.xml'):
            if name in names:
                root=etree.fromstring(archive.read(name))
                for note in root:
                    value=''.join(note.itertext()).strip()
                    if value:
                        layout.append({'kind':'p','block':block(value,'脚注/尾注')})
                warnings.append('脚注/尾注已整理至正文末尾。')
        archive.close()
        warnings.append('Word 以干净文档重新生成，保留文字、表格和插图；复杂排版、修订、批注和链接不保留。')
        return {'blocks':blocks,'layout':layout,'warnings':list(dict.fromkeys(warnings))}
    if extension=='xlsx':
        from openpyxl import load_workbook
        package(data,extension).close()
        workbook=load_workbook(BytesIO(data),data_only=False,keep_links=False)
        cached=load_workbook(BytesIO(data),data_only=True,keep_links=False)
        sheets=[]
        total=0
        for ws in workbook:
            if ws.sheet_state!='visible':
                warnings.append('隐藏工作表已排除，不写入成品。')
                continue
            if ws._charts:
                raise UserError('Excel 含图表，请先将图表转为静态图片或将文件转换为 PDF。')
            cells=[]
            for cell in list(ws._cells.values()):
                if cell.value is None:
                    continue
                from openpyxl.utils.cell import column_index_from_string
                hidden_column=any(d.hidden and (d.min or column_index_from_string(key))<=cell.column<=(d.max or column_index_from_string(key)) for key,d in ws.column_dimensions.items())
                if ws.row_dimensions[cell.row].hidden or hidden_column:
                    warnings.append('隐藏行列已排除，不写入成品。')
                    continue
                total+=1
                if total>200000 or cell.row>100000 or cell.column>1000:
                    raise UserError('Excel 使用区域超过本版本上限，请拆分工作簿。')
                value=cell.value
                if cell.data_type=='f':
                    value=cached[ws.title][cell.coordinate].value
                    if value is None:
                        raise UserError('Excel 公式没有可靠的缓存值，请在 Excel 中计算并保存后重新上传。')
                    warnings.append('公式已固定为保存时的静态值，不保留公式或外链。')
                if hasattr(value,'isoformat'):
                    value=value.isoformat(sep=' ') if hasattr(value,'hour') else value.isoformat()
                context=str(ws.cell(1,cell.column).value or '') if cell.row>1 else ''
                ref=block(str(value),f'工作表 {len(sheets)+1} / {cell.coordinate}',context)
                safe_format=cell.number_format if re.fullmatch(r'[0#.,%/ :hHmsSyYdD\-@]+',cell.number_format) else 'General'
                cells.append({'row':cell.row,'col':cell.column,'block':ref,'value':value,'format':safe_format})
            images=[]
            for item in ws._images:
                image=load_image(item._data())
                images.append({'image':add_image(image),'anchor':getattr(item.anchor,'_from',None).row+1 if getattr(item.anchor,'_from',None) else 1})
            sheets.append({'name':block(ws.title,'工作表名称'),'cells':cells,'images':images,
                           'widths':{key:dimension.width for key,dimension in ws.column_dimensions.items() if dimension.width and not dimension.hidden}})
        workbook.close()
        cached.close()
        if not sheets:
            raise UserError('工作簿中没有可见工作表。')
        warnings.append('Excel 输出可编辑静态值；移除批注、链接、隐藏内容与原有打印范围。')
        return {'blocks':blocks,'layout':sheets,'warnings':list(dict.fromkeys(warnings))}
    raise UserError('该格式不使用文本解析器。')


def native_export(document, extension, texts, read_image):
    out=BytesIO()
    changed=0
    if extension=='txt':
        return ''.join(texts[item['block']] for item in document['layout']).encode('utf-8'),0
    if extension=='csv':
        stream=StringIO(newline='')
        writer=csv.writer(stream,delimiter=document['delimiter'])
        for row in document['layout']:
            safe=[csv_safe(texts[ref]) for ref in row]
            changed+=sum(v[1] for v in safe)
            writer.writerow([v[0] for v in safe])
        return stream.getvalue().encode('utf-8-sig'),changed
    if extension=='docx':
        from docx import Document
        from docx.shared import Inches
        doc=Document()
        def add_nodes(parent,nodes):
            for item in nodes:
                if item['kind']=='p':
                    parent.add_paragraph(texts[item['block']])
                elif item['kind']=='image':
                    parent.add_paragraph().add_run().add_picture(BytesIO(read_image(item['image'])),width=Inches(5.8 if parent is doc else 1.5))
                elif item['kind']=='table' and item['rows']:
                    width=max(len(row) for row in item['rows'])
                    table=parent.add_table(rows=len(item['rows']),cols=width) if parent is doc else parent.add_table(rows=len(item['rows']),cols=width,width=Inches(5.8))
                    table.style='Table Grid'
                    for row_index,row in enumerate(item['rows']):
                        for column_index,cell in enumerate(row):
                            add_nodes(table.cell(row_index,column_index),cell)
        add_nodes(doc,document['layout'])
        for attr in ('author','last_modified_by','title','subject','keywords','comments','category','identifier'):
            setattr(doc.core_properties,attr,'')
        doc.save(out)
        return out.getvalue(),0
    if extension=='xlsx':
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment
        workbook=Workbook()
        workbook.remove(workbook.active)
        for index,sheet in enumerate(document['layout']):
            title=re.sub(r'[\\/*?:\[\]]','_',texts[sheet['name']])[:25] or '工作表'
            ws=workbook.create_sheet(f'{title}_{index+1}')
            for key,width in sheet['widths'].items():
                ws.column_dimensions[key].width=min(width,80)
            for item in sheet['cells']:
                text=texts[item['block']]
                value=item['value'] if text==str(item['value']) else text
                cell=ws.cell(item['row'],item['col'],value)
                if isinstance(value,str):
                    cell.data_type='s'
                cell.number_format=item['format'] if not isinstance(value,str) else '@'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            for image in sheet['images']:
                picture=XLImage(BytesIO(read_image(image['image'])))
                ws.add_image(picture,f'A{image["anchor"]}')
        workbook.properties.creator=''
        workbook.properties.lastModifiedBy=''
        workbook.save(out)
        workbook.close()
        return out.getvalue(),0
    raise UserError('无法生成指定格式。')


def text_pages(document, extension):
    """Explicit simplified previews: wrapping, no silent row or column truncation."""
    records=[]
    if extension=='csv':
        for row in document['layout']:
            records.extend(row)
    elif extension=='xlsx':
        for sheet in document['layout']:
            records.append(sheet['name'])
            records.extend(item['block'] for item in sheet['cells'])
    else:
        def walk(nodes):
            for item in nodes:
                if item['kind']=='p':
                    records.append(item['block'])
                elif item['kind']=='table':
                    for row in item['rows']:
                        for cell in row:
                            walk(cell)
        walk(document['layout'])
    lookup={item['id']:item for item in document['blocks']}
    font_path=Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts/msyh.ttc'
    font=ImageFont.truetype(str(font_path),24)
    image=Image.new('RGB',(1240,1754),'white')
    draw=ImageDraw.Draw(image)
    y=55
    refs=[]
    for ref in records:
        text=lookup[ref]['text']
        value=(lookup[ref]['label']+'  '+text).replace('\r\n','\n').replace('\r','\n')
        line=''
        for character in value+'\n':
            if character=='\n' or draw.textlength(line+character,font=font)>1120:
                if y>1670:
                    yield image,list(dict.fromkeys(refs))
                    image=Image.new('RGB',(1240,1754),'white');draw=ImageDraw.Draw(image);y=55;refs=[]
                draw.text((55,y),line,font=font,fill='#243042')
                refs.append(ref)
                y+=34
                line='' if character=='\n' else character
            else:
                line+=character
        y+=12
    if refs or not records:
        yield image,list(dict.fromkeys(refs))
