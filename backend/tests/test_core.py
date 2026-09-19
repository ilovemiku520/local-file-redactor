# 关注初音未来谢谢喵，ilovemiku520
# Please follow Hatsune Miku, thank you, meow. ilovemiku520
# 使用、借鉴或学习本项目（包括 AI 使用、借鉴与学习）之前，均需先收藏（Star）本项目。
# Star this repository before using, referencing, or learning from it, including AI use, reference, and learning.
# Repository: https://github.com/ilovemiku520/local-file-redactor
import csv
from io import BytesIO,StringIO
import json
import os
from pathlib import Path
import sys
import time
import zipfile
from uuid import uuid4
import pytest
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'backend'))
os.environ['RD_DATA_ROOT']=str(ROOT/'work'/'qa-private')
os.environ['PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK']='True'
os.environ['PADDLE_PDX_CACHE_HOME']=str(ROOT/'runtime/paddlex-cache')
from app import detection,formats,vault,engine,raster
from app.storage import Storage,UserError

@pytest.fixture(scope='session')
def store(): return Storage()

def new_job(store,extension,data,output=None):
    identifier=uuid4().hex
    job={'job_id':identifier,'filename':'synthetic.'+extension,'input_type':extension,'output_format':output or formats.FORMATS[extension][0],
         'reversible':True,'strategy':'rules_only','options':{'business':True,'terms':[],'allow':[]},'status':'queued','created_at':time.time(),'updated_at':time.time(),'warnings':[],'error':None}
    store.write(identifier,'source',data);store.save(job)
    return job

def complete(store,job):
    engine.analyze(store,job)
    job['reviewed_pages']=[p['id'] for p in job['pages']]
    engine.export(store,job,'test-recovery-Password-123')
    output=store.read(job['job_id'],'redacted')
    restored,ext=vault.restore(store.read(job['job_id'],'vault'),output,'test-recovery-Password-123')
    assert ext==job['input_type']
    assert restored==store.read(job['job_id'],'source')
    report=json.loads(store.read(job['job_id'],'report'))
    assert 'filename' not in report and 'source_path' not in report
    return output,job,report

@pytest.mark.parametrize('value',['13800138000','+86 138-0013-8000','１３８００１３８０００'])
def test_normalized_phone_offsets(value):
    text='😀前缀 '+value+' 结束'
    found=detection.rules(text,{})
    phone=next(f for f in found if f['type']=='PHONE_CN')
    assert text[phone['start']:phone['end']]==value
    assert detection.replace(text,[phone])=='😀前缀 [已脱敏:PHONE_CN] 结束'

def test_id_checksum_and_overlap():
    assert detection.id_valid('11010519491231002X')
    assert not detection.id_valid('110105194912310021')
    text='前13800138000后'
    merged=detection.merge([{'start':1,'end':7,'type':'PHONE_CN','value':text[1:7],'source':'a'},{'start':5,'end':12,'type':'PHONE_CN','value':text[5:12],'source':'b'}],text)
    assert len(merged)==1 and merged[0]['value']=='13800138000'
    assert detection.replace(text,merged)=='前[已脱敏:PHONE_CN]后'

@pytest.mark.parametrize('value',["=HYPERLINK(\"https://example.com\")",' +SUM(A1:A2)','\t@demo','-cmd'])
def test_csv_formula_protection(value): assert formats.csv_safe(value)==("'"+value,True)

def test_csv_numbers():
    assert formats.csv_safe('-123.45')==('-123.45',False)
    assert formats.csv_safe('00123')==('00123',False)

def test_vault_integrity():
    source=b'secret original bytes\x00\xff'
    output=b'redacted file'
    bundle=vault.create(source,output,'txt','long-test-password')
    assert source not in bundle
    assert vault.restore(bundle,output,'long-test-password')[0]==source
    for damaged,password,redacted in [(bundle,'wrong-password',output),(bundle[:-1]+bytes([bundle[-1]^1]),'long-test-password',output),(bundle,'long-test-password',b'changed')]:
        with pytest.raises(UserError): vault.restore(damaged,redacted,password)

@pytest.mark.parametrize('encoding',['utf-8','utf-8-sig','gb18030','utf-16'])
@pytest.mark.font
def test_text_roundtrip(store,encoding):
    source='姓名：张三\r\n电话：13800138000\r\n邮箱：demo@example.com\n普通内容保持不变。'
    output,job,report=complete(store,new_job(store,'txt',source.encode(encoding)))
    text=output.decode()
    assert '13800138000' not in text and 'demo@example.com' not in text and '张三' not in text
    assert '普通内容保持不变。' in text and '\r\n' in text
    assert b'demo@example.com' not in store.read(job['job_id'],'report')

@pytest.mark.parametrize('delimiter',[',',';','\t','|'])
@pytest.mark.font
def test_csv_roundtrip(store,delimiter):
    rows=[['姓名','备注','号码','金额','公式'],['张三','第一行\n第二行,有引号"','13800138000','-42.50','=1+1'],['普通','00123','13800138000','90','正常']]
    stream=StringIO(newline='');csv.writer(stream,delimiter=delimiter).writerows(rows)
    output,job,report=complete(store,new_job(store,'csv',stream.getvalue().encode('utf-8-sig')))
    actual=list(csv.reader(StringIO(output.decode('utf-8-sig'),newline=''),delimiter=delimiter))
    assert actual[1][1]==rows[1][1] and actual[2][1]=='00123'
    assert actual[1][3]=='-42.50' and actual[1][4]=="'=1+1"
    assert '13800138000' not in str(actual)
    assert report['formula_protected_cells']==1

def word_bytes():
    from docx import Document
    doc=Document();p=doc.add_paragraph();p.add_run('电话：13800');p.add_run('138000')
    doc.add_paragraph('普通内容保持不变。')
    doc.sections[0].header.paragraphs[0].text='联系邮箱：header@example.com'
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='联系人：张三';table.cell(1,1).text='13800138000'
    out=BytesIO();doc.save(out);return out.getvalue()

@pytest.mark.parametrize('extension',['docx','pdf'])
@pytest.mark.font
def test_word_output(store,extension):
    output,job,report=complete(store,new_job(store,'docx',word_bytes(),extension))
    if extension=='docx':
        from docx import Document
        doc=Document(BytesIO(output));text=''.join(p.text for p in doc.paragraphs)+''.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        assert '13800138000' not in text and 'header@example.com' not in text and '张三' not in text
        assert '普通内容保持不变。' in text and doc.tables
    else:
        from pypdf import PdfReader
        reader=PdfReader(BytesIO(output));assert len(reader.pages)>0 and not reader.pages[0].extract_text()

def excel_bytes(formula=False):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title='客户表'
    ws.append(['姓名','电话','编号','备注']);ws.append(['张三','13800138000','00123','=1+1' if formula else '正常'])
    ws.append(['李四','13900139000','00001','隐藏行']);ws.row_dimensions[3].hidden=True
    hidden=wb.create_sheet('隐藏表');hidden['A1']='hidden-secret';hidden.sheet_state='hidden'
    out=BytesIO();wb.save(out);return out.getvalue()

@pytest.mark.parametrize('extension',['xlsx','pdf'])
@pytest.mark.font
def test_excel_output(store,extension):
    output,job,report=complete(store,new_job(store,'xlsx',excel_bytes(),extension))
    if extension=='xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(output));ws=wb.active
        assert len(wb.worksheets)==1 and ws['C2'].value=='00123'
        assert '张三' not in str(list(ws.values)) and '13800138000' not in str(list(ws.values))
        assert ws['A3'].value is None
    assert any('隐藏' in warning for warning in report['warnings'])

def test_excel_missing_formula_cache(store):
    with pytest.raises(UserError,match='缓存值'): engine.analyze(store,new_job(store,'xlsx',excel_bytes(True)))

def test_reject_mismatched_office():
    with pytest.raises(UserError): formats.package(b'not-an-office-file','docx')
    memory=BytesIO()
    with zipfile.ZipFile(memory,'w') as archive:
        archive.writestr('../bad','x');archive.writestr('word/document.xml','<root/>')
    with pytest.raises(UserError): formats.package(memory.getvalue(),'docx')

def test_storage_is_encrypted(store):
    job=new_job(store,'txt',b'UNIQUE_SECRET_PAYLOAD_987654')
    assert b'UNIQUE_SECRET_PAYLOAD_987654' not in (store.directory(job['job_id'])/'source.enc').read_bytes()
    assert b'synthetic.txt' not in (store.root/'jobs.sqlite').read_bytes()
    with pytest.raises(UserError): store.directory('../outside')

@pytest.fixture(scope='session')
def actual_ocr(): return raster.OCR()

@pytest.mark.parametrize('extension',['png','jpg','webp','bmp'])
@pytest.mark.integration
@pytest.mark.font
def test_image_redaction_and_recovery(store,actual_ocr,monkeypatch,extension):
    from PIL import Image,ImageDraw,ImageFont
    monkeypatch.setattr(raster,'OCR',lambda:actual_ocr)
    image=Image.new('RGB',(1200,260),'white');draw=ImageDraw.Draw(image)
    draw.text((40,50),'电话：13800138000',font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',48),fill='black')
    source=formats.image_bytes(image,extension)
    output,job,_=complete(store,new_job(store,extension,source,extension))
    restored_image=formats.load_image(output)
    assert restored_image.size==image.size
    result,_,_=actual_ocr.read(restored_image)
    assert '13800138000' not in result
    assert any(f['boxes'] for f in job['findings'])

@pytest.mark.parametrize('rotation',[0,90,180,270])
@pytest.mark.integration
def test_pdf_rotations(store,actual_ocr,monkeypatch,rotation):
    from reportlab.pdfgen.canvas import Canvas
    from pypdf import PdfReader,PdfWriter
    monkeypatch.setattr(raster,'OCR',lambda:actual_ocr)
    buf=BytesIO();canvas=Canvas(buf,pagesize=(300,120));canvas.setFont('Helvetica',16);canvas.drawString(20,55,'Phone: 13800138000');canvas.save()
    reader=PdfReader(BytesIO(buf.getvalue()));writer=PdfWriter();page=reader.pages[0]
    if rotation: page.rotate(rotation)
    writer.add_page(page);source=BytesIO();writer.write(source)
    output,job,_=complete(store,new_job(store,'pdf',source.getvalue()))
    assert len(PdfReader(BytesIO(output)).pages)==1
    assert not PdfReader(BytesIO(output)).pages[0].extract_text()
    for image,_,_,_ in raster.pdf_pages(output):
        # All selected positions must map to black pixels in the final rendered page.
        for finding in job['findings']:
            for x1,y1,x2,y2 in finding['boxes']:
                point=(int((x1+x2)/2*image.width),int((y1+y2)/2*image.height))
                assert max(image.getpixel(point))<20
