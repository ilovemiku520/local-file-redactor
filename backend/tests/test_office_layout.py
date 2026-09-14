"""Behavioral regression: native styling survives while source content does not."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import zipfile

from lxml import etree as ET
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import detection, formats, office, vault
from app.storage import UserError


def save(document):
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def styled_word():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.1)
    section.left_margin, section.right_margin = Cm(2.8), Cm(2.2)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    title = doc.add_paragraph('合同信息表', 'Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(.74)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.5
    for value, name, size, bold, color in [('电话：', '宋体', 12, True, '145DA0'),
                                          ('13800', 'Calibri', 12, False, '8B0000'),
                                          ('138000', 'Courier New', 13, True, '8B0000'),
                                          ('  普通说明保持不变。', '微软雅黑', 11, False, '333333')]:
        run = p.add_run(value)
        run.font.name, run.font.size, run.bold = name, Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)
    p.add_run().add_tab()
    p.add_run('尾部')
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Light Shading Accent 1'
    table.autofit = False
    table.cell(0, 0).merge(table.cell(0, 2)).text = '保留合并表头'
    table.cell(1, 0).text = '姓名：张三'
    table.cell(1, 1).text = '普通项目'
    table.cell(2, 2).text = '4500.00'
    table.rows[1].height = Cm(1.1)
    section.header.paragraphs[0].text = '联系邮箱：header@example.com'
    section.header.paragraphs[0].runs[0].italic = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    doc.add_paragraph('第二页说明').paragraph_format.page_break_before = True
    return save(doc)


def styled_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from datetime import datetime
    wb = Workbook()
    ws = wb.active
    ws.title = '客户登记'
    ws.merge_cells('A1:D1')
    ws['A1'] = '客户登记表'
    ws['A1'].font = Font(name='微软雅黑', size=20, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor='145DA0')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.append(['联系人', '电话', '金额', '日期'])
    ws['A3'] = '姓名：张三'
    ws['B3'] = CellRichText(TextBlock(InlineFont(rFont='Calibri', b=True, sz=12), '13800'),
                            TextBlock(InlineFont(rFont='Courier New', i=True, sz=12), '138000'))
    ws['C3'], ws['D3'] = 1234.5, datetime(2026, 9, 14)
    ws['C3'].number_format = '#,##0.00'
    ws['D3'].number_format = 'yyyy-mm-dd'
    for row in ws.iter_rows(min_row=2, max_row=3, max_col=4):
        for cell in row:
            cell.font = Font(name='宋体', size=11)
            cell.border = Border(bottom=Side(style='thin', color='AABBCC'))
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for key, width in [('A', 22), ('B', 25), ('C', 18), ('D', 20)]:
        ws.column_dimensions[key].width = width
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[3].height = 31
    ws.freeze_panes = 'C3'
    ws.sheet_view.zoomScale = 85
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = 'A1:D6'
    ws.print_title_rows = '1:2'
    ws.oddHeader.center.text = '电话：13800138000'
    ws.oddHeader.center.font = 'Calibri,Bold'
    ws.oddHeader.center.size = 12
    ws.oddFooter.right.text = '第 &P 页 / 共 &N 页'
    ws['A5'] = 'hidden-secret-value'
    ws.row_dimensions[5].hidden = True
    hidden = wb.create_sheet('隐藏资料')
    hidden['A1'] = 'hidden-sheet-secret'
    hidden.sheet_state = 'hidden'
    return save(wb)


def exported(source, kind, terms=None):
    images = {}
    def add(image):
        index = len(images)
        images[index] = formats.image_bytes(image)
        return index
    doc = formats.parse(source, kind, add)
    findings = {b['id']: detection.merge(detection.rules(b['text'], {'terms': terms or [], 'business': True}, b.get('context', '')), b['text']) for b in doc['blocks']}
    texts = {b['id']: office.mask_text(b['text'], findings[b['id']]) for b in doc['blocks']}
    data, _ = formats.native_export(doc, kind, texts, images.__getitem__, original=source)
    return data, doc


def xml_parts(data):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return {name: ET.fromstring(archive.read(name)) for name in archive.namelist() if name.endswith(('.xml', '.rels'))}


def canonical(node):
    if node is None:
        return None
    node = deepcopy(node)
    for descendant in node.iter():
        for attr in list(descendant.attrib):
            if ET.QName(attr).localname.startswith('rsid'):
                del descendant.attrib[attr]
    return ET.tostring(node, method='c14n')


def property_values(data, names):
    return {path: [canonical(node) for node in root.iter() if ET.QName(node).localname in names]
            for path, root in xml_parts(data).items() if any(ET.QName(n).localname in names for n in root.iter())}


def rewrite_zip(data, changes):
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as old, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as new:
        for name in old.namelist():
            if name not in changes:
                new.writestr(name, old.read(name))
        for name, value in changes.items():
            new.writestr(name, value)
    return out.getvalue()


def test_word_preserves_runs_sections_tables_headers_and_breaks():
    from docx import Document
    source = styled_word()
    result, _ = exported(source, 'docx')
    # Compare actual Word property structures, not the implementation's plan.
    properties = {'pPr', 'rPr', 'sectPr', 'tblPr', 'tblGrid', 'tcPr', 'trPr'}
    assert property_values(source, properties) == property_values(result, properties)
    doc = Document(BytesIO(result))
    assert doc.paragraphs[1].runs[1].text == '*****'
    assert doc.paragraphs[1].runs[2].text == '******'
    assert '普通说明保持不变' in doc.paragraphs[1].text
    assert doc.sections[0].header.paragraphs[0].text == '联系邮箱：******************'
    assert len(doc.tables[0].rows[0].cells) == 3
    assert '■' in doc.tables[0].cell(1, 0).text
    assert len(doc.paragraphs) == len(Document(BytesIO(source)).paragraphs)
    bundle = vault.create(source, result, 'docx', 'FormatTestPassword123')
    assert vault.restore(bundle, result, 'FormatTestPassword123')[0] == source


def test_excel_preserves_rich_runs_styles_merges_print_layout_and_types():
    from openpyxl import load_workbook
    source = styled_excel()
    result, _ = exported(source, 'xlsx')
    props = {'fonts', 'fills', 'borders', 'cellXfs', 'numFmts', 'sheetViews', 'cols',
             'mergeCells', 'pageMargins', 'pageSetup', 'printOptions', 'sheetPr', 'definedNames'}
    before_properties = property_values(source, props)
    before_properties.pop('xl/worksheets/sheet2.xml')  # Deliberately excluded hidden sheet.
    assert before_properties == property_values(result, props)
    before = load_workbook(BytesIO(source), rich_text=True)
    after = load_workbook(BytesIO(result), rich_text=True)
    ws = after.active
    assert ws.title == '客户登记' and len(after.worksheets) == 1
    assert str(ws['B3'].value) == '*' * 11
    assert [r.font for r in ws['B3'].value] == [r.font for r in before.active['B3'].value]
    for address in ('A1', 'A3', 'B3', 'C3', 'D3'):
        assert ws[address]._style == before.active[address]._style
    assert ws['C3'].value == 1234.5 and ws['D3'].value == before.active['D3'].value
    assert ws.row_dimensions[3].height == 31 and ws.row_dimensions[5].hidden
    assert ws['A5'].value is None
    assert ws.oddHeader.center.font == before.active.oddHeader.center.font
    assert ws.oddHeader.center.size == 12
    assert '13800138000' not in ws.oddHeader.center.text
    assert ws.oddFooter.right.text == before.active.oddFooter.right.text
    assert all('hidden-secret' not in ET.tostring(r).decode() for r in xml_parts(result).values())
    assert 'hidden-sheet-secret' not in str(list(ws.values))


def test_word_textbox_and_footnotes_keep_their_original_locations():
    from docx import Document
    source = styled_word()
    parts = xml_parts(source)
    root = parts['word/document.xml']
    p = ET.SubElement(root.find(office.q('w', 'body')), office.q('w', 'p'))
    run = ET.SubElement(p, office.q('w', 'r'))
    box = ET.SubElement(run, office.q('w', 'txbxContent'))
    inner = ET.SubElement(box, office.q('w', 'p'))
    ET.SubElement(ET.SubElement(inner, office.q('w', 'r')), office.q('w', 't')).text = '电话：13900139000'
    note = ET.Element(office.q('w', 'footnotes'), nsmap={'w': office.NS['w']})
    item = ET.SubElement(note, office.q('w', 'footnote'), {office.q('w', 'id'): '1'})
    ET.SubElement(ET.SubElement(ET.SubElement(item, office.q('w', 'p')), office.q('w', 'r')), office.q('w', 't')).text = '邮箱：note@example.com'
    rels = parts['word/_rels/document.xml.rels']
    ET.SubElement(rels, office.q('p', 'Relationship'), Id='rIdNotes', Type=office.REL_BASE + 'footnotes', Target='footnotes.xml')
    types = parts['[Content_Types].xml']
    ET.SubElement(types, office.q('ct', 'Override'), PartName='/word/footnotes.xml', ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    source = rewrite_zip(source, {'word/document.xml': office.encoded(root), 'word/footnotes.xml': office.encoded(note),
                                  'word/_rels/document.xml.rels': office.encoded(rels), '[Content_Types].xml': office.encoded(types)})
    result, doc = exported(source, 'docx')
    assert sum(b['text'].count('13900139000') for b in doc['blocks']) == 1
    output = xml_parts(result)
    assert list(output['word/document.xml'].iter(office.q('w', 'txbxContent')))
    assert 'note@example.com' not in ''.join(output['word/footnotes.xml'].itertext())
    assert len(Document(BytesIO(result)).paragraphs) == len(Document(BytesIO(source)).paragraphs)


def test_word_removes_revisions_metadata_hidden_content_and_external_targets():
    source = styled_word()
    parts = xml_parts(source)
    root = parts['word/document.xml']
    body = root.find(office.q('w', 'body'))
    p = ET.SubElement(body, office.q('w', 'p'))
    deleted = ET.SubElement(p, office.q('w', 'del'), {office.q('w', 'author'): 'private-author'})
    ET.SubElement(ET.SubElement(deleted, office.q('w', 'r')), office.q('w', 'delText')).text = 'deleted-secret'
    hidden = ET.SubElement(p, office.q('w', 'r'))
    ET.SubElement(ET.SubElement(hidden, office.q('w', 'rPr')), office.q('w', 'vanish'))
    ET.SubElement(hidden, office.q('w', 't')).text = 'hidden-secret'
    hyperlink = ET.SubElement(p, office.q('w', 'hyperlink'), {office.q('r', 'id'): 'rIdPrivate'})
    ET.SubElement(ET.SubElement(hyperlink, office.q('w', 'r')), office.q('w', 't')).text = '普通链接文字'
    rels = parts['word/_rels/document.xml.rels']
    ET.SubElement(rels, office.q('p', 'Relationship'), Id='rIdPrivate', Type=office.REL_BASE + 'hyperlink', Target='https://example.com/private-secret', TargetMode='External')
    source = rewrite_zip(source, {'word/document.xml': office.encoded(root), 'word/_rels/document.xml.rels': office.encoded(rels),
                                  'customXml/private.xml': b'<data>custom-secret</data>', 'word/comments.xml': b'<data>comment-secret</data>'})
    result, _ = exported(source, 'docx')
    visible = ''.join(ET.tostring(r).decode() for r in xml_parts(result).values())
    for secret in ('deleted-secret', 'hidden-secret', 'private-secret', 'private-author', 'custom-secret', 'comment-secret'):
        assert secret not in visible
    assert '普通链接文字' in ''.join(xml_parts(result)['word/document.xml'].itertext())
    with zipfile.ZipFile(BytesIO(result)) as archive:
        assert not any(n.startswith(('customXml/', 'docProps/')) or 'comments' in n for n in archive.namelist())


def test_shared_strings_redacted_per_cell_and_unused_entries_removed():
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    wb.active['A1'], wb.active['B1'] = 'placeholder', 'placeholder'
    source = save(wb)
    parts = xml_parts(source)
    sheet = parts['xl/worksheets/sheet1.xml']
    for cell in sheet.iter(office.q('s', 'c')):
        for child in list(cell):
            cell.remove(child)
        cell.set('t', 's')
        ET.SubElement(cell, office.q('s', 'v')).text = '0'
    shared = ET.Element(office.q('s', 'sst'), nsmap={None: office.NS['s']})
    for value in ('13800138000', 'unused-private-secret'):
        ET.SubElement(ET.SubElement(shared, office.q('s', 'si')), office.q('s', 't')).text = value
    rels = parts['xl/_rels/workbook.xml.rels']
    ET.SubElement(rels, office.q('p', 'Relationship'), Id='rIdShared', Type=office.REL_BASE + 'sharedStrings', Target='sharedStrings.xml')
    source = rewrite_zip(source, {'xl/worksheets/sheet1.xml': office.encoded(sheet), 'xl/sharedStrings.xml': office.encoded(shared), 'xl/_rels/workbook.xml.rels': office.encoded(rels)})
    doc = formats.parse(source, 'xlsx', lambda _: 0)
    texts = {b['id']: b['text'] for b in doc['blocks']}
    first = doc['layout'][0]['cells'][0]['block']
    texts[first] = '*' * 11
    result, _ = formats.native_export(doc, 'xlsx', texts, lambda _: b'', original=source)
    ws = load_workbook(BytesIO(result)).active
    assert ws['A1'].value == '*' * 11 and ws['B1'].value == '13800138000'
    with zipfile.ZipFile(BytesIO(result)) as archive:
        assert 'xl/sharedStrings.xml' not in archive.namelist()


def test_excel_formula_cache_becomes_static_with_original_number_format():
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    wb.active['A1'] = '=1+1'
    wb.active['A1'].number_format = '0.00'
    source = save(wb)
    root = xml_parts(source)['xl/worksheets/sheet1.xml']
    root.find('.//' + office.q('s', 'v')).text = '2'
    source = rewrite_zip(source, {'xl/worksheets/sheet1.xml': office.encoded(root)})
    result, _ = exported(source, 'xlsx')
    cell = load_workbook(BytesIO(result)).active['A1']
    assert cell.value == 2 and cell.data_type == 'n' and cell.number_format == '0.00'


def test_image_anchor_size_and_pixels_survive_package_rewrite():
    from docx import Document
    from docx.shared import Cm
    from PIL import Image
    image = Image.new('RGB', (120, 60), '#336699')
    doc = Document()
    doc.add_picture(BytesIO(formats.image_bytes(image)), width=Cm(3))
    source = save(doc)
    result, _ = exported(source, 'docx')
    after = Document(BytesIO(result))
    assert after.inline_shapes[0].width == Document(BytesIO(source)).inline_shapes[0].width
    with zipfile.ZipFile(BytesIO(result)) as archive:
        media = [n for n in archive.namelist() if n.startswith('word/media/')]
        assert len(media) == 1 and media[0].endswith('.png')
        assert Image.open(BytesIO(archive.read(media[0]))).getpixel((40, 20)) == (51, 102, 153)


def test_old_task_requires_reanalysis_instead_of_silent_layout_loss():
    with pytest.raises(UserError, match='旧版排版'):
        formats.native_export({'layout': [], 'blocks': []}, 'docx', {}, None, original=styled_word())


@pytest.mark.integration
@pytest.mark.font
def test_styled_word_with_real_ocr_export_and_recovery(tmp_path, monkeypatch):
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from docx.shared import Cm
    from app import engine, raster
    from app.storage import Storage
    from test_core import new_job, complete
    ocr = raster.OCR()
    monkeypatch.setattr(raster, 'OCR', lambda: ocr)
    image = Image.new('RGB', (1000, 180), 'white')
    ImageDraw.Draw(image).text((30, 40), '电话：13800138000', font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 42), fill='black')
    doc = Document(BytesIO(styled_word()))
    doc.add_picture(BytesIO(formats.image_bytes(image)), width=Cm(12))
    source = save(doc)
    store = Storage(tmp_path)
    result, job, report = complete(store, new_job(store, 'docx', source))
    assert report['layout_mode'] == 'ooxml_preserved'
    assert any(f['boxes'] for f in job['findings'])
    assert Document(BytesIO(result)).inline_shapes[0].width == Document(BytesIO(source)).inline_shapes[0].width
    with zipfile.ZipFile(BytesIO(result)) as archive:
        media = [n for n in archive.namelist() if n.startswith('word/media/')]
        assert len(media) == 1
        text, _, _ = ocr.read(Image.open(BytesIO(archive.read(media[0]))))
        assert '13800138000' not in text
